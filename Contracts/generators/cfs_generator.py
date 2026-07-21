from __future__ import annotations

import datetime
import logging
import tempfile
import zipfile
from calendar import monthrange
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from docxtpl import DocxTemplate

from Contracts.shared.file_utils import create_zip_from_bytes, sanitize_filename
from Contracts.shared.pdf_utils import convert_docx_to_pdf


PROJECT_ROOT = Path(__file__).resolve().parents[2]
CFS_TEMPLATE_PATH = PROJECT_ROOT / "Contracts" / "templates" / "CFS" / "AMS - CFS - REB - Template.docx"
LOGGER = logging.getLogger(__name__)

CFS_MAIN_SECTION_HEADINGS = (
    "Appointment",
    "Independent Contractor Relationship",
    "Term",
    "Scope of Services",
    "Engagement Arrangement",
    "Fees and Payment Terms",
    "Expenses and Reimbursements",
    "Confidentiality",
    "Non-Solicitation",
    "Liability and Indemnity",
    "Vehicle Use and Damage",
    "Termination",
    "Intellectual Property",
    "Conflict of Interest",
    "Data Protection",
    "Dispute Resolution",
    "Governing Law",
    "Entire Agreement",
)
CFS_ANNEX_SECTION_HEADINGS = (
    "1. Service Responsibilities",
    "2. Operational Requirements",
    "3. Safety and Compliance",
    "4. Conduct",
    "5. Reporting of Incidents",
    "6. Service Standards",
)

# Printable writing lines used when producing an unfilled copy of the contract.
# Underscore characters are intentional: they remain visible in both Word and
# printed copies, unlike underlined whitespace which Word may collapse.
BLANK_CFS_CONTEXT = {
    "agreement_date": "____________________",
    "contractor_name": "________________________________________",
    "nric": "____________________",
    "residential_address": "____________________________________________________________",
    "start_date": "____________________",
    "end_date": "____________________",
    "service_start_time": "____________",
    "service_end_time": "____________",
    "service_fee": "____________",
}


@dataclass(frozen=True)
class BulkContractFailure:
    """One row that could not be rendered or converted during bulk generation."""

    row_number: int
    identifier: str
    exception_type: str
    message: str

    def as_dict(self) -> dict[str, str | int]:
        return {
            "Row Number": self.row_number,
            "Contractor": self.identifier,
            "Error Type": self.exception_type,
            "Issue": self.message,
        }


@dataclass(frozen=True)
class BulkContractResult:
    """Successful PDF archive and failures from an independent-row bulk run."""

    zip_bytes: bytes | None
    generated_filenames: tuple[str, ...]
    failures: tuple[BulkContractFailure, ...]

    @property
    def successful_count(self) -> int:
        return len(self.generated_filenames)


def end_of_month(value: datetime.date) -> datetime.date:
    """Return the final calendar day in the month containing ``value``."""
    return value.replace(day=monthrange(value.year, value.month)[1])


def format_contract_date(value: datetime.date) -> str:
    """Format date as e.g. '30 June 2026' or '1 July 2026' (no leading zeros)."""
    if not value:
        return ""
    return f"{value.day} {value.strftime('%B %Y')}"


def format_contract_time(value: datetime.time) -> str:
    """Format time as e.g. '2:00 p.m.' or '5:00 p.m.' matching legal document styles."""
    if not value:
        return ""
    hour = str(value.hour % 12 or 12)
    minute = value.strftime("%M")
    ampm = "a.m." if value.hour < 12 else "p.m."
    return f"{hour}:{minute} {ampm}"


def build_contract_context(
    agreement_date: datetime.date,
    contractor_name: str,
    nric: str,
    residential_address: str,
    start_date: datetime.date,
    end_date: datetime.date,
    service_start_time: datetime.time,
    service_end_time: datetime.time,
    service_fee: float,
) -> dict:
    """Build the dictionary of values to render into the contract template."""
    return {
        "agreement_date": format_contract_date(agreement_date),
        "contractor_name": contractor_name.strip().upper(),
        "nric": nric.strip().upper(),
        "residential_address": residential_address.strip(),
        "start_date": format_contract_date(start_date),
        "end_date": format_contract_date(end_date),
        "service_start_time": format_contract_time(service_start_time),
        "service_end_time": format_contract_time(service_end_time),
        "service_fee": f"{service_fee:.2f}",
    }


def _normalise_paragraph_text(value: str) -> str:
    return " ".join(value.split())


def _keep_paragraph_group_together(paragraphs, start_index: int, end_index: int) -> None:
    """Keep a logical clause together when it fits, while allowing oversized clauses to flow."""
    meaningful_indices = [
        index
        for index in range(start_index, end_index)
        if _normalise_paragraph_text(paragraphs[index].text)
    ]
    if not meaningful_indices:
        return

    last_meaningful_index = meaningful_indices[-1]
    for index in range(start_index, last_meaningful_index):
        paragraphs[index].paragraph_format.keep_with_next = True


def _remove_manual_page_breaks(paragraph) -> None:
    """Replace embedded page-break runs with a stable paragraph pagination property."""
    for page_break in paragraph._p.xpath(".//w:br[@w:type='page']"):
        page_break.getparent().remove(page_break)


def _apply_cfs_pagination(document) -> None:
    """Apply legal-document pagination rules to a rendered CFS contract."""
    paragraphs = document.paragraphs
    normalised_text = [_normalise_paragraph_text(paragraph.text) for paragraph in paragraphs]

    for paragraph in paragraphs:
        paragraph_format = paragraph.paragraph_format
        paragraph_format.keep_together = True
        paragraph_format.keep_with_next = False
        paragraph_format.widow_control = True

    main_heading_indices = [
        index for index, text in enumerate(normalised_text) if text in CFS_MAIN_SECTION_HEADINGS
    ]
    signature_index = next(
        (index for index, text in enumerate(normalised_text) if text == "SIGNATURES"),
        len(paragraphs),
    )
    annex_index = next(
        (index for index, text in enumerate(normalised_text) if text == "Annex A – Scope of Services"),
        len(paragraphs),
    )

    if main_heading_indices:
        _keep_paragraph_group_together(paragraphs, 0, main_heading_indices[0])
        main_group_boundaries = main_heading_indices[1:] + [signature_index]
        for start_index, end_index in zip(main_heading_indices, main_group_boundaries):
            _keep_paragraph_group_together(paragraphs, start_index, end_index)

        entire_agreement_index = next(
            (
                index
                for index in main_heading_indices
                if normalised_text[index] == "Entire Agreement"
            ),
            None,
        )
        if entire_agreement_index is not None:
            paragraphs[entire_agreement_index].paragraph_format.page_break_before = True

    if signature_index < annex_index:
        _keep_paragraph_group_together(paragraphs, signature_index, annex_index)

    if annex_index < len(paragraphs):
        annex_paragraph = paragraphs[annex_index]
        _remove_manual_page_breaks(annex_paragraph)
        annex_paragraph.paragraph_format.page_break_before = True

        annex_heading_indices = [
            index
            for index, text in enumerate(normalised_text)
            if index > annex_index and text in CFS_ANNEX_SECTION_HEADINGS
        ]
        if annex_heading_indices:
            _keep_paragraph_group_together(paragraphs, annex_index, annex_heading_indices[0])
            annex_group_boundaries = annex_heading_indices[1:] + [len(paragraphs)]
            for start_index, end_index in zip(annex_heading_indices, annex_group_boundaries):
                _keep_paragraph_group_together(paragraphs, start_index, end_index)


def generate_cfs_docx(context: dict, template_path: Path = CFS_TEMPLATE_PATH) -> BytesIO:
    """Render the CFS Word template and return the file bytes in memory."""
    if not template_path.exists():
        raise FileNotFoundError("The base contract template file could not be found.")

    template = DocxTemplate(str(template_path))
    template.render(context)
    rendered_template = BytesIO()
    template.save(rendered_template)
    rendered_template.seek(0)

    document = Document(rendered_template)
    _apply_cfs_pagination(document)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def generate_blank_cfs_docx(template_path: Path = CFS_TEMPLATE_PATH) -> BytesIO:
    """Return a printable CFS form with writing lines in every fill-in field."""
    return generate_cfs_docx(BLANK_CFS_CONTEXT, template_path)


def ensure_no_unresolved_placeholders(docx_bytes: bytes) -> None:
    """Fail fast if rendered DOCX XML still contains Jinja markers."""
    with zipfile.ZipFile(BytesIO(docx_bytes)) as docx_zip:
        for name in docx_zip.namelist():
            if not name.endswith(".xml"):
                continue
            xml = docx_zip.read(name).decode("utf-8", errors="ignore")
            if "{{" in xml or "{%" in xml or "{#" in xml:
                raise RuntimeError("Rendered contract still contains unresolved template placeholders.")


def generate_cfs_pdf(data: dict, output_path: Path | None = None) -> bytes:
    """Render one CFS contract and return PDF bytes."""
    with tempfile.TemporaryDirectory(prefix="ams_cfs_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        target_pdf_path = Path(output_path).resolve() if output_path is not None else temp_dir / "contract.pdf"
        docx_path = temp_dir / f"{target_pdf_path.stem}.docx"
        docx_bytes = generate_cfs_docx(data).getvalue()
        ensure_no_unresolved_placeholders(docx_bytes)
        docx_path.write_bytes(docx_bytes)
        pdf_path = convert_docx_to_pdf(docx_path, target_pdf_path.parent)
        return pdf_path.read_bytes()


def _unique_archive_filename(filename: str, used_names: set[str]) -> str:
    """Keep the legacy name unless a prior row already uses it in the ZIP."""
    candidate = filename
    suffix = Path(filename).suffix
    stem = Path(filename).stem
    counter = 2
    while candidate.casefold() in used_names:
        candidate = f"{stem} ({counter}){suffix}"
        counter += 1
    used_names.add(candidate.casefold())
    return candidate


def _display_row_number(row_index: object, position: int) -> int:
    try:
        return int(row_index) + 1
    except (TypeError, ValueError, OverflowError):
        return position


def build_bulk_contract_batch(
    contractors: pd.DataFrame,
    agreement_date: datetime.date,
    start_date: datetime.date,
    end_date: datetime.date,
    service_start_time: datetime.time,
    service_end_time: datetime.time,
    service_fee: float,
    progress,
) -> BulkContractResult:
    """Generate all valid CFS rows independently and retain every successful PDF."""
    total = len(contractors)
    generated_files: list[tuple[str, bytes]] = []
    generated_filenames: list[str] = []
    failures: list[BulkContractFailure] = []
    used_archive_names: set[str] = set()

    with tempfile.TemporaryDirectory(prefix="ams_contracts_") as operation_dir_name:
        for position, (row_index, contractor) in enumerate(contractors.iterrows(), start=1):
            if progress is not None:
                progress.progress(
                    (position - 1) / total,
                    text=f"Generating contract {position} of {total}",
                )

            row_number = _display_row_number(row_index, position)
            identifier = str(contractor.get("Full Name", "") or contractor.get("NRIC", "") or "Unknown")

            try:
                with tempfile.TemporaryDirectory(
                    prefix=f"row_{position}_",
                    dir=operation_dir_name,
                ) as row_dir_name:
                    row_dir = Path(row_dir_name)
                    context = build_contract_context(
                        agreement_date=agreement_date,
                        contractor_name=contractor["Full Name"],
                        nric=contractor["NRIC"],
                        residential_address=contractor["Residential Address"],
                        start_date=start_date,
                        end_date=end_date,
                        service_start_time=service_start_time,
                        service_end_time=service_end_time,
                        service_fee=service_fee,
                    )
                    docx_bytes = generate_cfs_docx(context).getvalue()
                    ensure_no_unresolved_placeholders(docx_bytes)
                    docx_path = row_dir / f"contract_{position}.docx"
                    docx_path.write_bytes(docx_bytes)
                    pdf_path = convert_docx_to_pdf(docx_path, row_dir)
                    pdf_bytes = pdf_path.read_bytes()

                safe_name = sanitize_filename(contractor["Full Name"])
                archive_name = _unique_archive_filename(
                    f"AMS - CFS - REB - {safe_name}.pdf",
                    used_archive_names,
                )
                generated_files.append((archive_name, pdf_bytes))
                generated_filenames.append(archive_name)
            except Exception as exc:
                LOGGER.exception(
                    "CFS bulk generation failed for displayed row %s (%s)",
                    row_number,
                    identifier,
                )
                failures.append(
                    BulkContractFailure(
                        row_number=row_number,
                        identifier=identifier,
                        exception_type=type(exc).__name__,
                        message=str(exc).strip() or "No error details were returned.",
                    )
                )

            if progress is not None:
                progress.progress(
                    position / total,
                    text=f"Generating contract {position} of {total}",
                )

    zip_bytes = create_zip_from_bytes(generated_files) if generated_files else None
    return BulkContractResult(
        zip_bytes=zip_bytes,
        generated_filenames=tuple(generated_filenames),
        failures=tuple(failures),
    )


def build_bulk_contract_zip(
    contractors: pd.DataFrame,
    agreement_date: datetime.date,
    start_date: datetime.date,
    end_date: datetime.date,
    service_start_time: datetime.time,
    service_end_time: datetime.time,
    service_fee: float,
    progress,
) -> bytes:
    """Compatibility wrapper returning the successful PDFs as the existing ZIP payload."""
    result = build_bulk_contract_batch(
        contractors=contractors,
        agreement_date=agreement_date,
        start_date=start_date,
        end_date=end_date,
        service_start_time=service_start_time,
        service_end_time=service_end_time,
        service_fee=service_fee,
        progress=progress,
    )
    if result.zip_bytes is None:
        if result.failures:
            first_failure = result.failures[0]
            raise RuntimeError(
                f"No PDF contracts were generated. Row {first_failure.row_number}: "
                f"{first_failure.exception_type}: {first_failure.message}"
            )
        raise RuntimeError("No PDF contracts were generated because the contractor list was empty.")
    return result.zip_bytes
